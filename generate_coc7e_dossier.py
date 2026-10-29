# -*- coding: utf-8 -*-
"""
Full regeneration script for generate_coc7e_dossier.py:
- Removes '冰涼小手' completely.
- Creates authentic 1920s Private Investigator: 艾德華·皮爾斯 (Edward Pierce).
- Exact 480-point allocation (STR:60, CON:55, SIZ:55, DEX:70, APP:45, INT:75, POW:60, EDU:60).
- Pure numeric age in cell C5 (fixes MOV = 9 calculation bug).
- Full DataValidation dropdowns on K4 (Method), K5 (127 Jobs from DB), D5 (Gender), E5 (Era), K15 (5F Acute Panic), Tab 2 Col C (Key Connection).
- Dual-mode checkbox support (TRUE / '☑') so formulas work seamlessly with native checkboxes AND dropdown toggles.
- Default hidden columns (Q:Z) and hidden rows (76+) on all worksheets.
"""

import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def build_coc7e_dossier(filename="COC7e_Flagship_Investigator_Dossier.xlsx"):
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    # Load Databases
    with open('occupations_clean.json', 'r', encoding='utf-8') as f:
        occupations = json.load(f)

    with open('weapons_clean.json', 'r', encoding='utf-8') as f:
        weapons = json.load(f)

    # Palette - 1920s Vintage Miskatonic Dossier
    C_MAHOGANY_DARK = "2C1A11"    # Dark Espresso Mahogany (Main Title)
    C_PARCHMENT_GOLD = "4A2E18"   # Rich Leather Header (Section Title)
    C_PARCHMENT_ACCENT = "8C6D46" # Old Gold Accent
    C_WARM_BEIGE = "EADCC9"       # Muted Antique Beige (Table Headers)
    C_LIGHT_CREAM = "F7F3EB"      # Aged Parchment Light Row
    C_CALC_TINT = "EFE8DC"        # Slightly darker parchment for calculated fields
    C_BORDER_INK = "6E543C"       # Sepia Ink Border
    C_TEXT_DARK = "211510"        # Charcoal Sepia Text
    C_TEXT_LIGHT = "FDFBF7"       # Antique Linen Light Text
    C_CR_GREEN = "2E4D34"         # Muted Forest Ink (Audit OK)

    font_main_banner = Font(name="Microsoft JhengHei", size=13, bold=True, color=C_TEXT_LIGHT)
    font_sub_banner = Font(name="Microsoft JhengHei", size=10, bold=True, color=C_TEXT_LIGHT)
    font_sec_header = Font(name="Microsoft JhengHei", size=10.5, bold=True, color=C_TEXT_LIGHT)
    font_tbl_header = Font(name="Microsoft JhengHei", size=9.5, bold=True, color=C_TEXT_DARK)
    font_bold_dark = Font(name="Microsoft JhengHei", size=9.5, bold=True, color=C_TEXT_DARK)
    font_regular = Font(name="Microsoft JhengHei", size=9.5, bold=False, color=C_TEXT_DARK)
    font_calc = Font(name="Microsoft JhengHei", size=9.5, bold=True, color=C_MAHOGANY_DARK)
    font_small_note = Font(name="Microsoft JhengHei", size=8.5, bold=False, italic=True, color="5A4535")

    fill_main_banner = PatternFill(start_color=C_MAHOGANY_DARK, end_color=C_MAHOGANY_DARK, fill_type="solid")
    fill_sec_header = PatternFill(start_color=C_PARCHMENT_GOLD, end_color=C_PARCHMENT_GOLD, fill_type="solid")
    fill_tbl_header = PatternFill(start_color=C_WARM_BEIGE, end_color=C_WARM_BEIGE, fill_type="solid")
    fill_light_row = PatternFill(start_color=C_LIGHT_CREAM, end_color=C_LIGHT_CREAM, fill_type="solid")
    fill_calc = PatternFill(start_color=C_CALC_TINT, end_color=C_CALC_TINT, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(style='thin', color=C_BORDER_INK)
    box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=False)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=False)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=False)

    def format_cell(ws, cell_ref, val, font, fill, align):
        cell = ws[cell_ref]
        if val is not None:
            cell.value = val
        if font: cell.font = font
        if fill is not None: cell.fill = fill
        if align: cell.alignment = align
        cell.border = box_border

    def format_range(ws, start_ref, end_ref, val, font, fill, align):
        ws.merge_cells(f"{start_ref}:{end_ref}")
        cell = ws[start_ref]
        if val is not None:
            cell.value = val
        cell.font = font
        if fill is not None: cell.fill = fill
        cell.alignment = align
        start_col, start_row = openpyxl.utils.coordinate_to_tuple(start_ref)
        end_col, end_row = openpyxl.utils.coordinate_to_tuple(end_ref)
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                cell_cur = ws.cell(row=r, column=c)
                if fill is not None: cell_cur.fill = fill
                cell_cur.border = box_border

    # =========================================================================
    # TAB 1: 📜 調查員核心檔案 (Core Dossier)
    # =========================================================================
    ws1 = wb.create_sheet(title="📜 調查員核心檔案")
    ws1.views.sheetView[0].showGridLines = True

    col_widths = {
        'A': 4.0,  # Checkbox Left
        'B': 22.0, # Skill Name Left
        'C': 6.5,  # Base / Age
        'D': 6.0,  # Occ / Gender
        'E': 6.0,  # Int / Era
        'F': 6.0,  # Growth / Education
        'G': 7.0,  # Total
        'H': 10.0, # Hard | Ext
        'I': 4.0,  # Checkbox Right
        'J': 22.0, # Skill Name Right
        'K': 6.5,  # Base / Current
        'L': 6.5,  # Occ / Max
        'M': 5.0,  # Int / Flag Checkbox
        'N': 9.0,  # Growth / Threshold
        'O': 7.5,  # Total
        'P': 10.5  # Hard | Ext
    }
    for col_letter, width in col_widths.items():
        ws1.column_dimensions[col_letter].width = width

    # Row 1: Main Header Banner
    ws1.merge_cells("A1:P1")
    c = ws1["A1"]
    c.value = "🏛️ MISKATONIC UNIVERSITY • DEPARTMENT OF SPECIAL INVESTIGATIONS"
    c.font = font_main_banner
    c.fill = fill_main_banner
    c.alignment = align_center
    ws1.row_dimensions[1].height = 28

    # Row 2: Sub-Banner
    ws1.merge_cells("A2:P2")
    c = ws1["A2"]
    c.value = "【 阿卡姆調查局 · 機密調查員檔案 (CALL OF CTHULHU 7th EDITION DOSSIER) 】"
    c.font = font_sub_banner
    c.fill = fill_sec_header
    c.alignment = align_center
    ws1.row_dimensions[2].height = 20

    # Row 3: Section Titles
    format_range(ws1, "A3", "H3", "📜 調查員基本身份檔案 (Investigator Profile)", font_sec_header, fill_sec_header, align_center)
    format_range(ws1, "I3", "P3", "📊 點數配置儀表板 (Skill & Attribute Dashboard)", font_sec_header, fill_sec_header, align_center)
    ws1.row_dimensions[3].height = 22

    # Row 4: Name, Player | Method, Stats Audit
    format_range(ws1, "A4", "B4", "姓名 (PC Name):", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "C4", "F4", "艾德華·皮爾斯 (Edward Pierce)", font_bold_dark, fill_white, align_left)
    format_range(ws1, "G4", "H4", "玩家: 調查員", font_bold_dark, fill_white, align_center)

    format_range(ws1, "I4", "J4", "創角模式 (Method):", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "K4", "M4", "480點分配法 (Companion)", font_bold_dark, fill_white, align_center)
    format_range(ws1, "N4", "P4", '=SUM(E10:E17)&" / 480 點 ("&IF(SUM(E10:E17)=480,"合規✔",IF(SUM(E10:E17)<480,"剩餘 "&(480-SUM(E10:E17))&" 點","超額 "&(SUM(E10:E17)-480)&" 點⚠"))&")"', font_calc, fill_calc, align_center)
    ws1.row_dimensions[4].height = 22

    # Row 5: Age (pure number in C5!), Gender, Era, Background | Job, Formula
    format_range(ws1, "A5", "B5", "年齡 / 性別 / 時代:", font_bold_dark, fill_tbl_header, align_right)
    format_cell(ws1, "C5", 34, font_bold_dark, fill_white, align_center) # PURE NUMBER AGE 34!
    format_cell(ws1, "D5", "男", font_regular, fill_white, align_center)
    format_cell(ws1, "E5", "1920s", font_regular, fill_white, align_center)
    format_range(ws1, "F5", "H5", "退伍軍人 / 私家偵探", font_regular, fill_white, align_center)

    format_range(ws1, "I5", "J5", "調查員職業 (Job):", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "K5", "M5", "私家偵探", font_bold_dark, fill_white, align_center)
    format_range(ws1, "N5", "P5", '=IFERROR(VLOOKUP(K5, \'📚 官方核心資料庫\'!C:F, 2, FALSE), "EDU*2+(STR or DEX)*2")', font_calc, fill_calc, align_center)
    ws1.row_dimensions[5].height = 22

    # Row 6: Birthplace & Residence | Points Pools
    format_range(ws1, "A6", "B6", "出生地 / 現居地:", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "C6", "H6", "麻薩諸塞州 波士頓  /  現居阿卡姆市 浸信會街42號 (皮爾斯偵探事務所)", font_regular, fill_white, align_left)

    format_range(ws1, "I6", "J6", "本職點 (總/餘):", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "K6", "L6", '=260 & "  /  剩: " & (260 - SUM(D22:D61) - SUM(L22:L61))', font_bold_dark, fill_calc, align_center)
    format_range(ws1, "M6", "N6", "個人興趣點 (總/餘):", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "O6", "P6", '=(E15*2) & "  /  剩: " & (E15*2 - SUM(E22:E61) - SUM(M22:M61))', font_bold_dark, fill_calc, align_center)
    ws1.row_dimensions[6].height = 22

    # Row 7: Appearance & CR
    format_range(ws1, "A7", "B7", "外貌特徵描述:", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "C7", "H7", "身材結實勻稱，神情警覺沉穩。常穿深灰色雙排扣風衣與呢帽，配戴皮革槍套與速記本。", font_regular, fill_white, align_left)

    format_range(ws1, "I7", "J7", "當前信用評級 (CR):", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "K7", "L7", '=G22 & " % (建議 " & IFERROR(VLOOKUP(K5, \'📚 官方核心資料庫\'!C:F, 3, FALSE), "9~30") & ")"', font_bold_dark, fill_calc, align_center)
    format_range(ws1, "M7", "P7", '=IF(G22=0, "請填寫CR ⚠", IF(AND(G22>=9, G22<=30), "CR合規 ✔ (普通中產階層)", IF(G22<9, "低於職業下限 ⚠", "高於職業上限 ⚠")))', font_bold_dark, fill_calc, align_center)
    ws1.row_dimensions[7].height = 22

    # Row 8: Section Headers
    format_range(ws1, "A8", "H8", "⚖️ 九大基礎特徵值 (Nine Characteristics - Core & Companion)", font_sec_header, fill_sec_header, align_center)
    format_range(ws1, "I8", "P8", "🩺 生理狀態與心智監控儀表 (Vitals & Sanity Monitor - Companion Revised)", font_sec_header, fill_sec_header, align_center)
    ws1.row_dimensions[8].height = 22

    # Row 9: Table Column Headers
    format_range(ws1, "A9", "B9", "特徵代號與名稱", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "C9", "基礎值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "D9", "年齡調整", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "E9", "最終值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "F9", "困難 1/2", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "G9", "極限 1/5", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "H9", "判定範圍", font_tbl_header, fill_tbl_header, align_center)

    format_range(ws1, "I9", "J9", "監控項目與指標", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "K9", "當前值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "L9", "上限", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "M9", "標記", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "N9", "臨界線", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws1, "O9", "P9", "即時綜合狀態判定 (純公式雙向連動)", font_tbl_header, fill_tbl_header, align_center)
    ws1.row_dimensions[9].height = 20

    # Rows 10-18: Nine Characteristics & Vitals Monitors
    # 480 Points Allocation (STR:60, CON:55, DEX:70, APP:45, POW:60, INT:75, SIZ:55, EDU:60 = 480 Exact!)
    chars_data = [
        ("STR", "力量 (Strength)", 60, 0),
        ("CON", "體質 (Constitution)", 55, 0),
        ("DEX", "敏捷 (Dexterity)", 70, 0),
        ("APP", "外貌 (Appearance)", 45, 0),
        ("POW", "意志 (Power)", 60, 0),
        ("INT", "智力/靈感 (Idea)", 75, 0),
        ("SIZ", "體型 (Size)", 55, 0),
        ("EDU", "教育 (Education)", 60, 0),
        ("LUCK", "幸運 (Luck)", 55, 0)
    ]

    for idx, (code, name, base_val, age_mod) in enumerate(chars_data):
        r = 10 + idx
        ws1.row_dimensions[r].height = 22
        format_range(ws1, f"A{r}", f"B{r}", f"{code} {name}", font_bold_dark, fill_light_row, align_left)
        format_cell(ws1, f"C{r}", base_val, font_regular, fill_white, align_center)
        format_cell(ws1, f"D{r}", age_mod, font_regular, fill_white, align_center)
        format_cell(ws1, f"E{r}", f"=C{r}+D{r}", font_bold_dark, fill_calc, align_center)
        if code != "LUCK":
            format_cell(ws1, f"F{r}", f"=INT(E{r}/2)", font_regular, fill_calc, align_center)
            format_cell(ws1, f"G{r}", f"=INT(E{r}/5)", font_regular, fill_calc, align_center)
            format_cell(ws1, f"H{r}", f'=E{r}&" ("&F{r}&"/"&G{r}&")"', font_calc, fill_calc, align_center)
        else:
            format_cell(ws1, f"F{r}", "-", font_regular, fill_calc, align_center)
            format_cell(ws1, f"G{r}", "-", font_regular, fill_calc, align_center)
            format_cell(ws1, f"H{r}", "無半值/極限", font_small_note, fill_calc, align_center)

    # Right Side: Vitals & Sanity Monitor (Rows 10-18)
    # Row 10: Hit Points
    format_range(ws1, "I10", "J10", "❤️ 生命值 (Hit Points)", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K10", 11, font_bold_dark, fill_white, align_center)
    format_cell(ws1, "L10", "=INT((E11+E16)/10)", font_calc, fill_calc, align_center)
    format_cell(ws1, "M10", None, font_bold_dark, fill_white, align_center) # Major Wound Checkbox
    format_cell(ws1, "N10", '="重傷: " & INT(L10/2) & "點"', font_bold_dark, fill_tbl_header, align_center)
    format_range(ws1, "O10", "P10", '=IF(K10<=-L10, "💀 瞬間致死 (Instant Death)", IF(AND(K10<=0, OR(M10=TRUE, M10="☑")), "⚠️ 瀕死 (Dying · 每輪CON檢定)", IF(K10<=0, "💤 昏迷 (Unconscious · 1小時甦醒)", IF(OR(M10=TRUE, M10="☑"), "🩹 重傷受創 (Major Wound · 休養每週恢復)", IF(K10<L10, "🤕 輕傷微恙 (Minor · 每日自然恢復1點)", "💚 健康健全 (Healthy)")))))', font_bold_dark, fill_calc, align_center)

    # Row 11: Magic Points
    format_range(ws1, "I11", "J11", "✨ 魔法值 (Magic Points)", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K11", 12, font_bold_dark, fill_white, align_center) # 60/5 = 12
    format_cell(ws1, "L11", "=INT(E14/5)", font_calc, fill_calc, align_center)
    format_cell(ws1, "M11", "-", font_regular, fill_calc, align_center)
    format_cell(ws1, "N11", "每小時+1點", font_small_note, fill_tbl_header, align_center)
    format_range(ws1, "O11", "P11", '=IF(K11<=0, "⚠️ 枯竭 (耗盡! 再消耗將扣減HP)", IF(K11<L11, "⏳ 充能中 (" & K11 & "/" & L11 & " · 自然恢復中)", "✨ 充盈健全 (滿額)"))', font_bold_dark, fill_calc, align_center)

    # Row 12: Sanity Points
    format_range(ws1, "I12", "J12", "🧠 當前理智 (Sanity)", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K12", 60, font_bold_dark, fill_white, align_center)
    format_cell(ws1, "L12", "=E14", font_calc, fill_calc, align_center)
    format_cell(ws1, "M12", None, font_bold_dark, fill_white, align_center)
    format_cell(ws1, "N12", '="上限: " & (99-O22)', font_bold_dark, fill_tbl_header, align_center)
    format_range(ws1, "O12", "P12", '=IF(K12<=0, "💀 永久瘋狂 (SAN歸零 · 角色退場)", IF(OR(M14=TRUE, M14="☑"), "🌀 心智質變中 (Changed · 幸運凍結)", IF(OR(M13=TRUE, M13="☑"), "⚡ 震懾餘波中 (Aftershock · 易生幻覺)", IF(K12<L12, "⚠️ 理智受損 (損耗 " & (L12-K12) & " 點)", "🧠 心智健全 (Sane)"))))', font_bold_dark, fill_calc, align_center)

    # Row 13: Shock
    format_range(ws1, "I13", "J13", "⚡ 震懾狀態 (Shock · 5F)", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K13", '=IF(OR(M14=TRUE, M14="☑"), 10, 5)', font_calc, fill_calc, align_center)
    format_cell(ws1, "L13", "1D10輪", font_small_note, fill_calc, align_center)
    format_cell(ws1, "M13", None, font_bold_dark, fill_white, align_center)
    format_cell(ws1, "N13", '="單次≥" & K13 & "點"', font_bold_dark, fill_tbl_header, align_center)
    format_range(ws1, "O13", "P13", '=IF(OR(M13=TRUE, M13="☑"), "⚡ 震懾餘波 (易受幻覺 · 可由安撫解除)", "心神平穩 (無震懾)")', font_bold_dark, fill_calc, align_center)

    # Row 14: Changed
    format_range(ws1, "I14", "J14", "🌀 心智質變 (Changed)", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K14", "=INT(E14/5)", font_calc, fill_calc, align_center)
    format_cell(ws1, "L14", "單日累積", font_small_note, fill_calc, align_center)
    format_cell(ws1, "M14", None, font_bold_dark, fill_white, align_center)
    format_cell(ws1, "N14", '="質變: " & INT(E14/5) & "點"', font_bold_dark, fill_tbl_header, align_center)
    format_range(ws1, "O14", "P14", '=IF(OR(M14=TRUE, M14="☑"), "🌀 心智質變 (幸運凍結 · 需改寫背景解質變)", "思想正常 (未質變)")', font_bold_dark, fill_calc, align_center)

    # Row 15: 5F Acute Panic Reaction
    format_range(ws1, "I15", "J15", "💥 5F 應激反應 (Acute Panic)", font_bold_dark, fill_light_row, align_left)
    format_range(ws1, "K15", "L15", "無 (健全)", font_bold_dark, fill_white, align_center)
    format_range(ws1, "M15", "N15", "持續 1D10 輪", font_small_note, fill_tbl_header, align_center)
    format_range(ws1, "O15", "P15", '=IF(AND(M13<>TRUE, M13<>"☑"), "心神平穩，無恐慌應激行為", IF(K15="討好 (Fawn)", "討好: 卑躬屈膝企圖安撫討好威脅者", IF(K15="戰鬥 (Fight)", "戰鬥: 狂怒襲擊眼前任何可疑事物", IF(K15="逃跑 (Flight)", "逃跑: 不顧一切瘋狂逃離恐懼現場", IF(K15="虛脫 (Flop)", "虛脫: 昏厥或全身虛脫無力癱倒在地", IF(K15="僵直 (Freeze)", "僵直: 目瞪口呆無法動彈、說話或思考", "處於震懾中，請在左欄選擇 5F 反應"))))))', font_regular, fill_calc, align_center)

    # Row 16: Damage Bonus & Build
    format_range(ws1, "I16", "J16", "🥊 傷害加值 (DB) / 體格", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K16", '=IF(E10+E16<65,"-2",IF(E10+E16<85,"-1",IF(E10+E16<125,"0",IF(E10+E16<165,"+1D4",IF(E10+E16<205,"+1D6",IF(E10+E16<285,"+2D6","+3D6"))))))', font_calc, fill_calc, align_center)
    format_cell(ws1, "L16", '=IF(E10+E16<65,-2,IF(E10+E16<85,-1,IF(E10+E16<125,0,IF(E10+E16<165,1,IF(E10+E16<205,2,IF(E10+E16<285,3,4))))))', font_bold_dark, fill_calc, align_center)
    format_range(ws1, "M16", "N16", '="STR+SIZ: " & (E10+E16)', font_bold_dark, fill_tbl_header, align_center)
    format_range(ws1, "O16", "P16", '=IF(L16<0,"身材矮小弱勢",IF(L16=0,"普通人標準體格",IF(L16=1,"魁梧健碩體格",IF(L16=2,"極其強壯彪悍","超乎常人之龐大"))))', font_small_note, fill_calc, align_center)

    # Row 17: Movement Rate & MvP (C5 is pure number 34, so age penalty is 0, MOV=9!)
    format_range(ws1, "I17", "J17", "🏃 移動力 (MOV) / 戰鬥 MvP", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K17", '=MAX(1, IF(AND(E10<E16, E12<E16), 7, IF(AND(E10>E16, E12>E16), 9, 8)) - IF(C5<40, 0, IF(C5<50, 1, IF(C5<60, 2, IF(C5<70, 3, IF(C5<80, 4, 5))))))', font_calc, fill_calc, align_center)
    format_cell(ws1, "L17", "=K17", font_bold_dark, fill_calc, align_center)
    format_range(ws1, "M17", "N17", "戰鬥行動點 (MvP)", font_bold_dark, fill_tbl_header, align_center)
    format_range(ws1, "O17", "P17", '="每輪可行進 " & (K17*5) & " 步 (或 " & K17 & " 碼)"', font_small_note, fill_calc, align_center)

    # Row 18: Reassurance & Key Connection
    format_range(ws1, "I18", "J18", "🤝 同伴安撫 / 精神錨點", font_bold_dark, fill_light_row, align_left)
    format_cell(ws1, "K18", "=INT(E13/5)", font_calc, fill_calc, align_center)
    format_cell(ws1, "L18", "本劇本限1次", font_small_note, fill_calc, align_center)
    format_range(ws1, "M18", "N18", "精神錨點休整", font_bold_dark, fill_tbl_header, align_center)
    format_range(ws1, "O18", "P18", "團間與精神錨點互動可恢復 1D6 SAN", font_small_note, fill_calc, align_center)

    # Row 19: Financial & Asset Dashboard
    ws1.row_dimensions[19].height = 22
    format_range(ws1, "A19", "B19", "💰 社會階層與經濟狀況:", font_bold_dark, fill_sec_header, align_right)
    format_range(ws1, "C19", "D19", "消費水準: $10 (免記帳)", font_regular, fill_calc, align_center)
    format_range(ws1, "E19", "F19", "隨身現金: $50", font_regular, fill_calc, align_center)
    format_range(ws1, "G19", "H19", "資產總額: $1,250", font_regular, fill_calc, align_center)
    format_range(ws1, "I19", "P19", "資產明細: 浸信會街二樓偵探事務所租約、老舊雪佛蘭四門轎車、阿卡姆商業銀行私人活存帳戶", font_regular, fill_white, align_left)

    # =========================================================================
    # SKILLS TREE (Rows 20-61: 40 Left Skills & 40 Right Skills with Custom Slots!)
    # =========================================================================
    ws1.row_dimensions[20].height = 22
    format_range(ws1, "A20", "H20", "📖 調查員官方技能樹 (左列 · 交際 / 運動 / 技術 / 戰鬥 / 自訂專精)", font_sec_header, fill_sec_header, align_center)
    format_range(ws1, "I20", "P20", "📖 調查員官方技能樹 (右列 · 感知 / 學術 / 科學 / 語言 / 神秘 / 自訂專精)", font_sec_header, fill_sec_header, align_center)

    ws1.row_dimensions[21].height = 20
    format_cell(ws1, "A21", "成長", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "B21", "技能名稱 (左列)", font_tbl_header, fill_tbl_header, align_left)
    format_cell(ws1, "C21", "初始", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "D21", "本職", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "E21", "興趣", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "F21", "成長", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "G21", "總值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "H21", "困難|極限", font_tbl_header, fill_tbl_header, align_center)

    format_cell(ws1, "I21", "成長", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "J21", "技能名稱 (右列)", font_tbl_header, fill_tbl_header, align_left)
    format_cell(ws1, "K21", "初始", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "L21", "本職", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "M21", "興趣", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "N21", "成長", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "O21", "總值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "P21", "困難|極限", font_tbl_header, fill_tbl_header, align_center)

    # Private Investigator Allocated Skills (260 Occ + 150 Int)
    left_skills = [
        ("信用評級 (Credit Rating)", 0, 25, 0, 0),        # Occ 25
        ("安撫 (Reassure)", "=INT(E13/5)", 0, 0, 0),
        ("魅惑 / 取悅 (Charm)", 15, 0, 0, 0),
        ("恐嚇 (Intimidate)", 15, 0, 0, 0),
        ("說服 (Persuade)", 10, 0, 0, 0),
        ("話術 / 欺騙 (Fast Talk)", 5, 35, 0, 0),         # Occ 35
        ("心理學 (Psychology)", 10, 50, 10, 0),           # Occ 50 + Int 10 = 70%
        ("攀爬 (Climb)", 20, 0, 0, 0),
        ("閃避 (Dodge)", "=INT(E12/2)", 0, 25, 0),         # Int 25 = 60%
        ("跳躍 (Jump)", 20, 0, 0, 0),
        ("隱密行動 / 潛行 (Stealth)", 20, 0, 20, 0),       # Int 20 = 40%
        ("游泳 (Swim)", 20, 0, 0, 0),
        ("投擲 (Throw)", 20, 0, 0, 0),
        ("騎術 (Ride)", 5, 0, 0, 0),
        ("汽車駕駛 (Drive Auto)", 20, 0, 20, 0),           # Int 20 = 40%
        ("駕駛（飛機 Aircraft）", 1, 0, 0, 0),
        ("駕駛（船隻 Boat）", 1, 0, 0, 0),
        ("駕駛（自訂載具）", 1, 0, 0, 0),
        ("電器維修 (Electrical Repair)", 10, 0, 0, 0),
        ("機械維修 (Mechanical Repair)", 10, 0, 0, 0),
        ("重機械操作 (Operate Heavy)", 1, 0, 0, 0),
        ("鎖匠 (Locksmith)", 1, 0, 0, 0),
        ("妙手 / 巧手 (Sleight of Hand)", 10, 0, 0, 0),
        ("催眠 (Hypnosis)", 1, 0, 0, 0),
        ("讀唇 (Read Lips)", 1, 0, 0, 0),
        ("喬裝 (Disguise)", 5, 20, 0, 0),                 # Occ 20
        ("格鬥（鬥毆 Brawl）", 25, 0, 35, 0),              # Int 35 = 60%
        ("格鬥（斧頭 Axe）", 15, 0, 0, 0),
        ("格鬥（劍術 Sword）", 20, 0, 0, 0),
        ("格鬥（連枷 Flail）", 15, 0, 0, 0),
        ("格鬥（矛 Spear）", 20, 0, 0, 0),
        ("格鬥（絞繩 Garrote）", 15, 0, 0, 0),
        ("格鬥（電鋸 Chainsaw）", 15, 0, 0, 0),
        ("格鬥（自訂專精 1）", 10, 0, 0, 0),
        ("格鬥（自訂專精 2）", 10, 0, 0, 0),
        ("射擊（手槍 Handgun）", 20, 0, 40, 0),            # Int 40 = 60%
        ("射擊（步槍/霰彈 Rifle）", 25, 0, 0, 0),
        ("射擊（衝鋒槍 SMG）", 15, 0, 0, 0),
        ("射擊（機槍 Machine Gun）", 15, 0, 0, 0),
        ("射擊（弓弩 Bow）", 15, 0, 0, 0)
    ]

    right_skills = [
        ("克蘇魯神話 (Cthulhu Mythos)", 0, 0, 0, 0),
        ("偵查 / 識破 (Spot Hidden)", 25, 45, 0, 0),      # Occ 45 = 70%
        ("聆聽 (Listen)", 20, 0, 0, 0),
        ("追蹤 (Track)", 10, 0, 0, 0),
        ("急救 (First Aid)", 30, 0, 0, 0),
        ("醫學 (Medicine)", 1, 0, 0, 0),
        ("精神分析 (Psychoanalysis)", 1, 0, 0, 0),
        ("導航 / 領航 (Navigate)", 10, 0, 0, 0),
        ("生存（荒野 Wilderness）", 10, 0, 0, 0),
        ("生存（自訂專精）", 10, 0, 0, 0),
        ("圖書館使用 (Library Use)", 20, 30, 0, 0),        # Occ 30 = 50%
        ("會計 (Accounting)", 5, 0, 0, 0),
        ("法律 (Law)", 5, 35, 0, 0),                      # Occ 35 = 40%
        ("歷史 (History)", 5, 0, 0, 0),
        ("考古學 (Archaeology)", 1, 0, 0, 0),
        ("人類學 (Anthropology)", 1, 0, 0, 0),
        ("宗教 (Religion)", 10, 0, 0, 0),
        ("神秘學 (Occult)", 5, 0, 0, 0),
        ("博物學 / 自然 (Natural World)", 10, 0, 0, 0),
        ("語言（母語 英語）", "=E17", 0, 0, 0),
        ("語言（外語一 拉丁語）", 1, 0, 0, 0),
        ("語言（外語二 法語）", 1, 0, 0, 0),
        ("語言（外語三 德語）", 1, 0, 0, 0),
        ("語言（自訂外語）", 1, 0, 0, 0),
        ("藝術與手藝（攝影 Photography）", 5, 20, 0, 0), # Occ 20 = 25%
        ("藝術與手藝（偽造 Forgery）", 5, 0, 0, 0),
        ("藝術與手藝（音樂/樂器）", 5, 0, 0, 0),
        ("藝術與手藝（自訂專精）", 5, 0, 0, 0),
        ("科學（化學 Chemistry）", 1, 0, 0, 0),
        ("科學（生物學 Biology）", 1, 0, 0, 0),
        ("科學（物理學 Physics）", 1, 0, 0, 0),
        ("科學（藥學 Pharmacy）", 1, 0, 0, 0),
        ("科學（地質學 Geology）", 1, 0, 0, 0),
        ("科學（天文學 Astronomy）", 1, 0, 0, 0),
        ("科學（法醫學 Forensics）", 1, 0, 0, 0),
        ("科學（數學 Mathematics）", 10, 0, 0, 0),
        ("科學（自訂專精）", 1, 0, 0, 0),
        ("【自訂技能 1】", 1, 0, 0, 0),
        ("【自訂技能 2】", 1, 0, 0, 0),
        ("【自訂技能 3】", 1, 0, 0, 0)
    ]

    for i in range(40):
        r = 22 + i
        ws1.row_dimensions[r].height = 19
        row_fill = fill_light_row if i % 2 == 1 else fill_white

        # Left Skill
        l_name, l_base, l_occ, l_int, l_gro = left_skills[i]
        format_cell(ws1, f"A{r}", None, font_regular, fill_white, align_center)
        format_cell(ws1, f"B{r}", l_name, font_regular, row_fill, align_left)
        format_cell(ws1, f"C{r}", l_base, font_regular, fill_calc, align_center)
        format_cell(ws1, f"D{r}", l_occ if l_occ > 0 else None, font_regular, fill_white, align_center)
        format_cell(ws1, f"E{r}", l_int if l_int > 0 else None, font_regular, fill_white, align_center)
        format_cell(ws1, f"F{r}", l_gro if l_gro > 0 else None, font_regular, fill_white, align_center)
        format_cell(ws1, f"G{r}", f"=C{r}+D{r}+E{r}+F{r}", font_bold_dark, fill_calc, align_center)
        format_cell(ws1, f"H{r}", f'=INT(G{r}/2)&" | "&INT(G{r}/5)', font_calc, fill_calc, align_center)

        # Right Skill
        r_name, r_base, r_occ, r_int, r_gro = right_skills[i]
        format_cell(ws1, f"I{r}", None, font_regular, fill_white, align_center)
        format_cell(ws1, f"J{r}", r_name, font_regular, row_fill, align_left)
        format_cell(ws1, f"K{r}", r_base, font_regular, fill_calc, align_center)
        format_cell(ws1, f"L{r}", r_occ if r_occ > 0 else None, font_regular, fill_white, align_center)
        format_cell(ws1, f"M{r}", r_int if r_int > 0 else None, font_regular, fill_white, align_center)
        format_cell(ws1, f"N{r}", r_gro if r_gro > 0 else None, font_regular, fill_white, align_center)
        format_cell(ws1, f"O{r}", f"=K{r}+L{r}+M{r}+N{r}", font_bold_dark, fill_calc, align_center)
        format_cell(ws1, f"P{r}", f'=INT(O{r}/2)&" | "&INT(O{r}/5)', font_calc, fill_calc, align_center)

    # =========================================================================
    # WEAPONS ARMORY & COMBAT MANEUVERS (Rows 62-75)
    # =========================================================================
    ws1.row_dimensions[62].height = 22
    format_range(ws1, "A62", "P62", "⚔️ 戰鬥武器軍械庫與戰術動作 (Weapons Armory & Combat Maneuvers)", font_sec_header, fill_sec_header, align_center)

    ws1.row_dimensions[63].height = 20
    format_range(ws1, "A63", "B63", "武器名稱 (Weapon)", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws1, "C63", "D63", "關聯技能 (Skill)", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "E63", "技能%", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws1, "F63", "G63", "基礎傷害 (Damage)", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "H63", "穿刺傷害", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws1, "I63", "J63", "射程/距離", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "K63", "每輪次數", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "L63", "彈匣/殘彈", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "M63", "故障值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws1, "N63", "耐久 HP", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws1, "O63", "P63", "武器特性/戰術備註", font_tbl_header, fill_tbl_header, align_center)

    weapons_data = [
        ("徒手格鬥 / 拳打腳踢", "格鬥（鬥毆）", "=G48", "1D3+DB", "-", "接觸", "1", "-", "-", "-", "可實施繳械、抱摔、擒抱戰技"),
        ("小型折疊匕首", "格鬥（鬥毆）", "=G48", "1D4+DB", "4+DB", "接觸", "1", "-", "-", "9", "穿刺冷兵器，極限成功造成穿刺"),
        (".38 口徑左輪手槍", "射擊（手槍）", "=G57", "1D10", "10+1D10", "15 米", "1(3)", "6 / 6", "100", "9", "可靠便攜，私家偵探隨身標配槍械"),
        ("12號霰彈槍 (雙管)", "射擊（步槍/霰彈）", "=G58", "4D6/2D6/1D6", "24/12/6", "10/20/50米", "1或2", "2 / 2", "100", "12", "面殺傷，近距離極高毀滅力"),
        ("重型木製手杖", "格鬥（鬥毆）", "=G48", "1D6+DB", "-", "接觸", "1", "-", "-", "15", "可偽裝成隨身手杖，堅固實用"),
        ("自訂備用武器 1", "自訂技能", 0, "-", "-", "-", "-", "-", "-", "-", "自由指定武器數據"),
        ("自訂備用武器 2", "自訂技能", 0, "-", "-", "-", "-", "-", "-", "-", "自由指定武器數據")
    ]

    for idx, (w_name, w_skill, w_pct, w_dmg, w_imp, w_rng, w_rate, w_ammo, w_mal, w_hp, w_note) in enumerate(weapons_data):
        r = 64 + idx
        ws1.row_dimensions[r].height = 19
        fill_w = fill_light_row if idx % 2 == 1 else fill_white
        format_range(ws1, f"A{r}", f"B{r}", w_name, font_bold_dark, fill_w, align_left)
        format_range(ws1, f"C{r}", f"D{r}", w_skill, font_regular, fill_w, align_center)
        format_cell(ws1, f"E{r}", w_pct, font_bold_dark, fill_calc, align_center)
        format_range(ws1, f"F{r}", f"G{r}", w_dmg, font_regular, fill_calc, align_center)
        format_cell(ws1, f"H{r}", w_imp, font_regular, fill_calc, align_center)
        format_range(ws1, f"I{r}", f"J{r}", w_rng, font_regular, fill_w, align_center)
        format_cell(ws1, f"K{r}", w_rate, font_regular, fill_w, align_center)
        format_cell(ws1, f"L{r}", w_ammo, font_regular, fill_white, align_center)
        format_cell(ws1, f"M{r}", w_mal, font_regular, fill_calc, align_center)
        format_cell(ws1, f"N{r}", w_hp, font_regular, fill_calc, align_center)
        format_range(ws1, f"O{r}", f"P{r}", w_note, font_small_note, fill_w, align_left)

    # Maneuver Comparator & Firearm Maneuvers (Rows 71-75)
    format_range(ws1, "A71", "H71", "🥋 戰術戰技體格比對器 (Fighting Maneuver Build Comparator)", font_sec_header, fill_sec_header, align_center)
    format_range(ws1, "I71", "P71", "🔫 槍械戰術動作速查 (Firearm Maneuvers - Companion p.68)", font_sec_header, fill_sec_header, align_center)
    ws1.row_dimensions[71].height = 20

    format_range(ws1, "A72", "B72", "我方體格 (My Build):", font_bold_dark, fill_tbl_header, align_right)
    format_cell(ws1, "C72", "=L16", font_bold_dark, fill_calc, align_center)
    format_range(ws1, "D72", "E72", "目標體格 (Target):", font_bold_dark, fill_tbl_header, align_right)
    format_cell(ws1, "F72", 1, font_bold_dark, fill_white, align_center)
    format_cell(ws1, "G72", "體格差距:", font_bold_dark, fill_tbl_header, align_right)
    format_cell(ws1, "H72", "=F72-C72", font_calc, fill_calc, align_center)

    format_range(ws1, "I72", "P72", '="近身射擊 (Point-Blank): 目標距離 ≤ DEX/5 呎 (約 " & INT(E12/5) & " 呎)，射擊自動獲得 +1 獎勵骰！"', font_regular, fill_white, align_left)
    ws1.row_dimensions[72].height = 20

    format_range(ws1, "A73", "B73", "戰技判定修正:", font_bold_dark, fill_tbl_header, align_right)
    format_range(ws1, "C73", "H73", '=IF(H72<=0, "正常對抗 (無懲罰骰)", IF(H72=1, "承受 1 顆懲罰骰 (-1P)", IF(H72=2, "承受 2 顆懲罰骰 (-2P)", "❌ 戰技無效！(體格差≥3無法摔倒巨獸)")))', font_bold_dark, fill_calc, align_center)

    format_range(ws1, "I73", "P73", "製造恐慌 (Cause Panic): 於公共場合鳴槍或亮槍，迫使人群潰散或目標進行意志/理智判定。", font_regular, fill_light_row, align_left)
    ws1.row_dimensions[73].height = 20

    format_range(ws1, "A74", "H74", "規則摘要: 戰技包含繳械、推倒、抱摔、擒抱。若目標體格大於我方3點以上則物理無法實施。", font_small_note, fill_light_row, align_left)
    format_range(ws1, "I74", "P74", "槍口命令 (Hold at Gunpoint): 持已準備槍械指著目標保留先攻，目標若試圖動作立即搶先射擊！", font_regular, fill_white, align_left)
    ws1.row_dimensions[74].height = 20

    format_range(ws1, "A75", "H75", "互通門檻: 任何格鬥專精不得低於最高格鬥之一半；任何射擊專精不得低於最高射擊之一半。", font_small_note, fill_light_row, align_left)
    format_range(ws1, "I75", "P75", "圍攻加成 (Outnumbered): 目標同一輪內反擊或閃避過一次後，後續所有近戰攻擊者獲 +1 獎勵骰！", font_regular, fill_light_row, align_left)
    ws1.row_dimensions[75].height = 20

    # DATA VALIDATION FOR WS1
    # 1. Method Dropdown (K4)
    dv_method = DataValidation(type="list", formula1='"480點分配法 (Companion),460點購點法 (Core),3D6隨機擲骰 (Standard),高超本領 (Pulp)"', allow_blank=True)
    ws1.add_data_validation(dv_method)
    dv_method.add("K4")

    # 2. Gender Dropdown (D5)
    dv_gender = DataValidation(type="list", formula1='"男,女,其他"', allow_blank=True)
    ws1.add_data_validation(dv_gender)
    dv_gender.add("D5")

    # 3. Era Dropdown (E5)
    dv_era = DataValidation(type="list", formula1='"1920s,1930s (Pulp),現代 (Modern),煤氣燈 (Gaslight),冷戰 (Cold War)"', allow_blank=True)
    ws1.add_data_validation(dv_era)
    dv_era.add("E5")

    # 4. Job Dropdown (K5)
    dv_job = DataValidation(type="list", formula1="='📚 官方核心資料庫'!$C$5:$C$131", allow_blank=True)
    ws1.add_data_validation(dv_job)
    dv_job.add("K5")

    # 5. 5F Panic Dropdown (K15)
    dv_5f = DataValidation(type="list", formula1='"無 (健全),討好 (Fawn),戰鬥 (Fight),逃跑 (Flight),虛脫 (Flop),僵直 (Freeze)"', allow_blank=True)
    ws1.add_data_validation(dv_5f)
    dv_5f.add("K15")

    # 6. Flag Checkboxes Dropdown/Validation for M10, M12, M13, M14
    dv_flags = DataValidation(type="list", formula1='"☑,☐"', allow_blank=True)
    ws1.add_data_validation(dv_flags)
    dv_flags.add("M10")
    dv_flags.add("M12")
    dv_flags.add("M13")
    dv_flags.add("M14")

    # 7. Skill Growth Checkboxes DataValidation for Col A & Col I (Rows 22-61)
    dv_growth = DataValidation(type="list", formula1='"☑,☐"', allow_blank=True)
    ws1.add_data_validation(dv_growth)
    for r in range(22, 62):
        dv_growth.add(f"A{r}")
        dv_growth.add(f"I{r}")

    # Hide unused columns Q to Z and rows 76 to 200 by default
    for c in range(17, 27):
        ws1.column_dimensions[get_column_letter(c)].hidden = True
    for r in range(76, 201):
        ws1.row_dimensions[r].hidden = True

    # =========================================================================
    # TAB 2: 🕯️ 調查員背景與心智印記 (Backstory & Mental Scars)
    # =========================================================================
    ws2 = wb.create_sheet(title="🕯️ 調查員背景與心智印記")
    ws2.views.sheetView[0].showGridLines = True

    ws2_widths = {
        'A': 5.0,  # No.
        'B': 22.0, # Element / Title
        'C': 12.0, # Anchor Checkbox
        'D': 25.0, # Content 1
        'E': 25.0, # Content 2
        'F': 25.0, # Content 3
        'G': 18.0, # Changed Revision Log
        'H': 18.0  # Notes
    }
    for col, w in ws2_widths.items():
        ws2.column_dimensions[col].width = w

    ws2.merge_cells("A1:H1")
    c = ws2["A1"]
    c.value = "🏛️ MISKATONIC ARCHIVES • INVESTIGATOR DOSSIER: BACKSTORY & MENTAL SCARS"
    c.font = font_main_banner
    c.fill = fill_main_banner
    c.alignment = align_center
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells("A2:H2")
    c = ws2["A2"]
    c.value = "【 調查員背景十要素、心靈關鍵連結與神話心智印記 (BACKSTORY & TRAUMA LOG) 】"
    c.font = font_sub_banner
    c.fill = fill_sec_header
    c.alignment = align_center
    ws2.row_dimensions[2].height = 20

    # Backstory Elements Headers
    format_range(ws2, "A3", "H3", "🕯️ 官方背景故事十要素與精神關鍵連結 (10 Official Backstory Elements & Key Connection)", font_sec_header, fill_sec_header, align_center)
    ws2.row_dimensions[3].height = 22

    format_cell(ws2, "A4", "序號", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "B4", "背景故事核心要素", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "C4", "精神關鍵連結?", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws2, "D4", "F4", "詳細描述內容 (角色設定與回憶)", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws2, "G4", "H4", "變異 / 質變改寫記錄 (Changed Rewrite)", font_tbl_header, fill_tbl_header, align_center)
    ws2.row_dimensions[4].height = 20

    backstory_data = [
        (1, "個人外貌特徵 (Personal Description)", "-", "身材結實勻稱，神情沉著警覺。常穿深灰色雙排扣風衣與呢帽，內襯平整，隨身攜帶皮革槍套與速記本。眼神銳利，帶著長期在暗處蹲點的深沉。"),
        (2, "意識形態與信念 (Ideology / Beliefs)", "-", "「真相隱藏在微小的細節與陰影深處，沒有任何偽裝能經受住嚴密的邏輯剖析。」"),
        (3, "重要之人 (Significant People)", "★ 精神錨點", "老搭檔 湯瑪斯·卡特 (Thomas Carter，前波士頓巡警，現於阿卡姆經營報攤，多次生死與共)。"),
        (4, "意義非凡之地 (Meaningful Locations)", "-", "阿卡姆浸信會街42號的皮爾斯私家偵探事務所，那是他在這座迷霧城市中唯一的避風港與檔案室。"),
        (5, "珍視之物 (Treasured Possessions)", "-", "一枚刻有一戰阿爾貢森林戰役編號的純銀軍用懷錶，指針永遠滴答走動，提醒自己時刻保持冷靜。"),
        (6, "個人特質 (Traits)", "-", "執著、敏銳、擅長捕捉他人說謊時的微表情，在極度危險下反而更顯冷靜。"),
        (7, "傷痕與疤痕 (Injuries & Scars)", "-", "右側肩膀有一道陳舊的流彈擦傷疤痕，每逢陰雨天會隱隱作痛。"),
        (8, "恐懼症與狂熱症 (Phobias & Manias)", "-", "無 (精神健常，但對過度幽閉且不見天日的深水水域抱有本能戒心)。"),
        (9, "神秘典籍與神器 (Arcane Tomes & Artifacts)", "-", "一本用油布包裹的老舊私人調查手記，裡面記錄著幾起阿卡姆無法結案的失蹤懸案。"),
        (10, "遭遇超自然實體紀錄 (Strange Entities)", "-", "曾在密斯卡托尼克河畔的廢棄倉庫外，於濃霧中目睹過非人的龐大水棲輪廓掠過。")
    ]

    for idx, (num, elem, anchor, desc) in enumerate(backstory_data):
        r = 5 + idx
        ws2.row_dimensions[r].height = 24
        fill_b = fill_light_row if idx % 2 == 1 else fill_white
        format_cell(ws2, f"A{r}", num, font_regular, fill_b, align_center)
        format_cell(ws2, f"B{r}", elem, font_bold_dark, fill_b, align_left)
        format_cell(ws2, f"C{r}", anchor, font_bold_dark, fill_calc if anchor != "-" else fill_b, align_center)
        format_range(ws2, f"D{r}", f"F{r}", desc, font_regular, fill_white, align_left)
        format_range(ws2, f"G{r}", f"H{r}", None, font_regular, fill_white, align_left)

    format_range(ws2, "A15", "H15", "💡 精神關鍵連結說明: 玩家在上方挑選一項作為精神錨點。團間休整與其互動可檢定回復 1D6 SAN；若錨點遭毀壞將直接扣減 1D6 SAN。\n💡 心智質變改寫: 當調查員陷入質變 (Changed) 狀態時，在團後休整期必須永久改寫上方其中一項背景要素以解除質變狀態。", font_small_note, fill_light_row, align_left)
    ws2.row_dimensions[15].height = 28

    # Trauma Log
    format_range(ws2, "A16", "H16", "⚡ 震懾 (Shock) 與 心智質變 (Changed) 歷程記錄簿 (Mental Trauma Log - Companion p.87-93)", font_sec_header, fill_sec_header, align_center)
    ws2.row_dimensions[16].height = 22

    format_cell(ws2, "A17", "日期/劇本", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws2, "B17", "C17", "遭遇恐怖事件 / 刺激源", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "D17", "SAN 損失", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "E17", "5F 震懾反應", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws2, "F17", "G17", "餘波幻覺 (Delusion) 表現", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "H17", "質變改寫處置", font_tbl_header, fill_tbl_header, align_center)
    ws2.row_dimensions[17].height = 20

    for i in range(10):
        r = 18 + i
        ws2.row_dimensions[r].height = 20
        fill_t = fill_light_row if i % 2 == 1 else fill_white
        format_cell(ws2, f"A{r}", f"第 {i+1} 幕", font_regular, fill_t, align_center)
        format_range(ws2, f"B{r}", f"C{r}", None, font_regular, fill_white, align_left)
        format_cell(ws2, f"D{r}", None, font_regular, fill_white, align_center)
        format_cell(ws2, f"E{r}", None, font_regular, fill_white, align_center)
        format_range(ws2, f"F{r}", f"G{r}", None, font_regular, fill_white, align_left)
        format_cell(ws2, f"H{r}", None, font_regular, fill_white, align_left)

    # Arcane Tomes & Spells (Rows 29-36)
    format_range(ws2, "A29", "H29", "📜 神秘學典籍研讀與習得法術清單 (Arcane Tomes & Learned Spells)", font_sec_header, fill_sec_header, align_center)
    ws2.row_dimensions[29].height = 22

    format_range(ws2, "A30", "B30", "法術/典籍名稱", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "C30", "類型", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "D30", "施法消耗 (MP/SAN)", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws2, "E30", "施法時間", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws2, "F30", "H30", "法術效果與限制描述", font_tbl_header, fill_tbl_header, align_center)
    ws2.row_dimensions[30].height = 20

    for i in range(6):
        r = 31 + i
        ws2.row_dimensions[r].height = 20
        fill_s = fill_light_row if i % 2 == 1 else fill_white
        format_range(ws2, f"A{r}", f"B{r}", None, font_regular, fill_white, align_left)
        format_cell(ws2, f"C{r}", None, font_regular, fill_s, align_center)
        format_cell(ws2, f"D{r}", None, font_regular, fill_white, align_center)
        format_cell(ws2, f"E{r}", None, font_regular, fill_s, align_center)
        format_range(ws2, f"F{r}", f"H{r}", None, font_regular, fill_white, align_left)

    # DataValidation for Tab 2 Key Connection (C5:C14)
    dv_anchor = DataValidation(type="list", formula1='"★ 精神錨點,-"', allow_blank=True)
    ws2.add_data_validation(dv_anchor)
    for r in range(5, 15):
        dv_anchor.add(f"C{r}")

    # Hide unused columns I to Z and rows 37 to 150 by default
    for c in range(9, 27):
        ws2.column_dimensions[get_column_letter(c)].hidden = True
    for r in range(37, 151):
        ws2.row_dimensions[r].hidden = True

    # =========================================================================
    # TAB 3: 📈 團後成長與經歷日誌 (Post-Game Growth Log)
    # =========================================================================
    ws3 = wb.create_sheet(title="📈 團後成長與經歷日誌")
    ws3.views.sheetView[0].showGridLines = True

    ws3_widths = {
        'A': 5.0,  # No.
        'B': 22.0, # Skill Name
        'C': 8.0,  # Current Value
        'D': 8.0,  # Roll 1D100
        'E': 10.0, # Check Result
        'F': 8.0,  # Gain 1D10
        'G': 8.0,  # New Value
        'H': 12.0, # Discord Log Ref
        'I': 5.0,  # Scenario ID
        'J': 20.0, # Scenario Title
        'K': 12.0, # Date / Keeper
        'L': 10.0, # SAN Award
        'M': 10.0, # Mythos Award
        'N': 10.0, # LUCK Award
        'O': 12.0, # Major Scars
        'P': 16.0  # Keeper Signature
    }
    for col, w in ws3_widths.items():
        ws3.column_dimensions[col].width = w

    ws3.merge_cells("A1:P1")
    c = ws3["A1"]
    c.value = "🏛️ MISKATONIC UNIVERSITY • INVESTIGATOR DOSSIER: GROWTH & CAMPAIGN LOG"
    c.font = font_main_banner
    c.fill = fill_main_banner
    c.alignment = align_center
    ws3.row_dimensions[1].height = 28

    ws3.merge_cells("A2:P2")
    c = ws3["A2"]
    c.value = "【 團後技能成長結算紀錄簿 與 模組冒險通關日誌 (INVESTIGATOR GROWTH ARCHIVES) 】"
    c.font = font_sub_banner
    c.fill = fill_sec_header
    c.alignment = align_center
    ws3.row_dimensions[2].height = 20

    # Section Headers
    format_range(ws3, "A3", "H3", "📈 團後技能成長檢定計算器 (Skill Improvement Rolls)", font_sec_header, fill_sec_header, align_center)
    format_range(ws3, "I3", "P3", "🗺️ 劇本通關歷程與經歷日誌 (Scenario Chronicles & Awards)", font_sec_header, fill_sec_header, align_center)
    ws3.row_dimensions[3].height = 22

    # Skill Growth Headers
    format_cell(ws3, "A4", "序號", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "B4", "待成長技能名稱", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "C4", "當前總值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "D4", "1D100檢定", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "E4", "成長判定", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "F4", "增長1D10", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "G4", "成長後新值", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "H4", "Discord紀錄", font_tbl_header, fill_tbl_header, align_center)

    # Scenario Log Headers
    format_cell(ws3, "I4", "序號", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "J4", "模組/劇本名稱", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "K4", "日期 / KP", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "L4", "SAN 獎懲", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "M4", "神話獎勵", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws3, "N4", "幸運獎勵", font_tbl_header, fill_tbl_header, align_center)
    format_range(ws3, "O4", "P4", "後續影響與調查員經歷簡記", font_tbl_header, fill_tbl_header, align_center)
    ws3.row_dimensions[4].height = 20

    # 15 Growth and Log Rows
    for i in range(15):
        r = 5 + i
        ws3.row_dimensions[r].height = 20
        fill_g = fill_light_row if i % 2 == 1 else fill_white

        # Left: Growth rows
        format_cell(ws3, f"A{r}", i + 1, font_regular, fill_g, align_center)
        format_cell(ws3, f"B{r}", None, font_bold_dark, fill_white, align_left)
        format_cell(ws3, f"C{r}", None, font_regular, fill_white, align_center)
        format_cell(ws3, f"D{r}", None, font_regular, fill_white, align_center)
        format_cell(ws3, f"E{r}", f'=IF(ISBLANK(D{r}), "", IF(OR(D{r}>C{r}, D{r}>=96), "增長成功 ✔", "未增長 ✘"))', font_bold_dark, fill_calc, align_center)
        format_cell(ws3, f"F{r}", None, font_bold_dark, fill_white, align_center)
        format_cell(ws3, f"G{r}", f'=IF(ISBLANK(C{r}), "", C{r} + IF(ISBLANK(F{r}), 0, F{r}))', font_bold_dark, fill_calc, align_center)
        format_cell(ws3, f"H{r}", None, font_small_note, fill_white, align_center)

        # Right: Log rows
        format_cell(ws3, f"I{r}", i + 1, font_regular, fill_g, align_center)
        format_cell(ws3, f"J{r}", None, font_bold_dark, fill_white, align_left)
        format_cell(ws3, f"K{r}", None, font_regular, fill_white, align_center)
        format_cell(ws3, f"L{r}", None, font_regular, fill_white, align_center)
        format_cell(ws3, f"M{r}", None, font_regular, fill_white, align_center)
        format_cell(ws3, f"N{r}", None, font_regular, fill_white, align_center)
        format_range(ws3, f"O{r}", f"P{r}", None, font_regular, fill_white, align_left)

    # Hide unused columns Q to Z and rows 58 to 150 by default
    for c in range(17, 27):
        ws3.column_dimensions[get_column_letter(c)].hidden = True
    for r in range(58, 151):
        ws3.row_dimensions[r].hidden = True

    # =========================================================================
    # TAB 4: 📚 官方核心資料庫 (Core Database)
    # =========================================================================
    ws4 = wb.create_sheet(title="📚 官方核心資料庫")
    ws4.views.sheetView[0].showGridLines = True

    ws4_widths = {
        'A': 6.0,  # ID
        'B': 14.0, # Category
        'C': 18.0, # Occupation Name
        'D': 24.0, # Formula
        'E': 14.0, # CR Range
        'F': 45.0, # Skills
        'G': 4.0,  # Separator
        'H': 6.0,  # Weapon ID
        'I': 14.0, # Weapon Type
        'J': 20.0, # Weapon Name
        'K': 18.0, # Weapon Skill
        'L': 14.0, # Damage
        'M': 12.0, # Impale
        'N': 14.0, # Range
        'O': 10.0, # Uses
        'P': 10.0, # Ammo
        'Q': 10.0  # Malfunction
    }
    for col, w in ws4_widths.items():
        ws4.column_dimensions[col].width = w

    ws4.merge_cells("A1:Q1")
    c = ws4["A1"]
    c.value = "🏛️ CALL OF CTHULHU 7th EDITION • OFFICIAL REFERENCE DATABASE"
    c.font = font_main_banner
    c.fill = fill_main_banner
    c.alignment = align_center
    ws4.row_dimensions[1].height = 28

    ws4.merge_cells("A2:Q2")
    c = ws4["A2"]
    c.value = "【 官方 127 種職業速查庫 與 105 種全套軍械庫數據清單 (OFFICIAL CODEX) 】"
    c.font = font_sub_banner
    c.fill = fill_sec_header
    c.alignment = align_center
    ws4.row_dimensions[2].height = 20

    format_range(ws4, "A3", "F3", "📚 官方 127 種調查員職業索引數據庫 (Occupations Database)", font_sec_header, fill_sec_header, align_center)
    format_range(ws4, "H3", "Q3", "⚔️ 官方 105 種完整戰鬥武器軍械裝備庫 (Weapons Database)", font_sec_header, fill_sec_header, align_center)
    ws4.row_dimensions[3].height = 22

    format_cell(ws4, "A4", "編號", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "B4", "分類", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "C4", "職業名稱", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "D4", "技能點數計算公式", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "E4", "信用評級區間", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "F4", "官方推薦本職技能清單", font_tbl_header, fill_tbl_header, align_center)

    format_cell(ws4, "H4", "編號", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "I4", "武器類別", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "J4", "武器名稱", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "K4", "關聯技能", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "L4", "基礎傷害", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "M4", "穿刺傷害", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "N4", "射程/距離", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "O4", "每輪次數", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "P4", "彈匣容量", font_tbl_header, fill_tbl_header, align_center)
    format_cell(ws4, "Q4", "故障值", font_tbl_header, fill_tbl_header, align_center)
    ws4.row_dimensions[4].height = 20

    for idx, occ in enumerate(occupations):
        r = 5 + idx
        ws4.row_dimensions[r].height = 19
        fill_o = fill_light_row if idx % 2 == 1 else fill_white
        format_cell(ws4, f"A{r}", idx + 1, font_regular, fill_o, align_center)
        format_cell(ws4, f"B{r}", occ.get('category', ''), font_regular, fill_o, align_center)
        format_cell(ws4, f"C{r}", occ.get('name', ''), font_bold_dark, fill_o, align_left)
        format_cell(ws4, f"D{r}", occ.get('formula', ''), font_calc, fill_calc, align_center)
        format_cell(ws4, f"E{r}", occ.get('cr', ''), font_bold_dark, fill_o, align_center)
        format_cell(ws4, f"F{r}", occ.get('skills', ''), font_regular, fill_o, align_left)

    for idx, w in enumerate(weapons):
        r = 5 + idx
        fill_w = fill_light_row if idx % 2 == 1 else fill_white
        format_cell(ws4, f"H{r}", idx + 1, font_regular, fill_w, align_center)
        format_cell(ws4, f"I{r}", w.get('category', ''), font_regular, fill_w, align_center)
        format_cell(ws4, f"J{r}", w.get('name', ''), font_bold_dark, fill_w, align_left)
        format_cell(ws4, f"K{r}", w.get('skill', ''), font_regular, fill_w, align_center)
        format_cell(ws4, f"L{r}", w.get('damage', ''), font_calc, fill_calc, align_center)
        format_cell(ws4, f"M{r}", w.get('impale', ''), font_calc, fill_calc, align_center)
        format_cell(ws4, f"N{r}", w.get('range', ''), font_regular, fill_w, align_center)
        format_cell(ws4, f"O{r}", w.get('uses_per_round', ''), font_regular, fill_w, align_center)
        format_cell(ws4, f"P{r}", w.get('capacity', ''), font_regular, fill_w, align_center)
        format_cell(ws4, f"Q{r}", w.get('malfunction', ''), font_regular, fill_calc, align_center)

    # Hide unused columns R to Z and rows 132 to 200 by default
    for c in range(18, 27):
        ws4.column_dimensions[get_column_letter(c)].hidden = True
    for r in range(132, 201):
        ws4.row_dimensions[r].hidden = True

    if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(default_sheet)

    wb.save(filename)
    print(f"Flagship Dossier regenerated successfully: {filename}")

if __name__ == "__main__":
    build_coc7e_dossier()
